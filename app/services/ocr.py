import gc
import logging
import os
import re
import html
import shutil
import tempfile
import threading
import time
from pathlib import Path
from typing import Optional, Dict, Any, List, Tuple

from .converter import BaseConverter
from ..utils.exceptions import OCRError
from .llm import LLMService

logger = logging.getLogger(__name__)

OCR_DPI = 200
IMAGE_EXTENSIONS = {'.png', '.jpg', '.jpeg', '.webp', '.bmp', '.tif', '.tiff'}

RAPID_QUALITIES = {
    'fast': {'model': 'tiny', 'dpi': 150, 'max_side': 2000},
    'standard': {'model': 'small', 'dpi': 200, 'max_side': 2000},
    'accurate': {'model': 'medium', 'dpi': 300, 'max_side': 4000},
}
DEFAULT_RAPID_QUALITY = 'accurate'
RAPID_SCRIPTS = ('latin', 'eslav', 'cyrillic', 'el', 'arabic', 'korean', 'th', 'devanagari', 'ta', 'te')
DEFAULT_RAPID_SCRIPT = 'latin'
LANG_TO_SCRIPT = {
    'ru': 'eslav', 'uk': 'eslav', 'be': 'eslav',
    'bg': 'cyrillic', 'sr': 'cyrillic', 'mk': 'cyrillic', 'kk': 'cyrillic', 'mn': 'cyrillic',
    'el': 'el', 'ar': 'arabic', 'fa': 'arabic', 'ur': 'arabic', 'ko': 'korean', 'th': 'th',
    'hi': 'devanagari', 'mr': 'devanagari', 'ne': 'devanagari', 'ta': 'ta', 'te': 'te',
}
DET_MIN_SIDE = 736


LAYOUT_QUALITY = 'standard'
SCRIPT_BLOCKS = (
    ('el', 0x0370, 0x03FF), ('cyr', 0x0400, 0x04FF), ('arabic', 0x0600, 0x06FF),
    ('devanagari', 0x0900, 0x097F), ('ta', 0x0B80, 0x0BFF), ('te', 0x0C00, 0x0C7F),
    ('th', 0x0E00, 0x0E7F), ('korean', 0xAC00, 0xD7AF),
)
MIN_ALIGN_RATIO = 0.4


class OcrPage:
    def __init__(self, text: str, lines: Optional[list] = None, layout: Optional[tuple] = None):
        self.text = text or ''
        self.lines = lines or []
        self.layout = layout


def _detect_script(text: str) -> str:
    counts: Dict[str, int] = {}
    for char in text:
        if not char.isalpha():
            continue
        code = ord(char)
        script = next((name for name, lo, hi in SCRIPT_BLOCKS if lo <= code <= hi), 'latin')
        counts[script] = counts.get(script, 0) + 1
    if not counts:
        return DEFAULT_RAPID_SCRIPT
    script = max(counts, key=counts.get)
    if script != 'cyr':
        return script
    lower = text.lower()
    if any(c in lower for c in 'ђјљњћџѓќѕ'):
        return 'cyrillic'
    if any(c in lower for c in 'ыэёіїєґў'):
        return 'eslav'
    return 'cyrillic' if 'ъ' in lower else 'eslav'


def _xy_cut(items) -> list:
    if len(items) <= 1:
        return list(items)
    best = None
    for lo, hi in ((1, 3), (0, 2)):
        ordered = sorted(items, key=lambda item: item[2][lo])
        reach = ordered[0][2][hi]
        for k in range(1, len(ordered)):
            gap = ordered[k][2][lo] - reach
            if gap > 0 and (best is None or gap > best[0]):
                best = (gap, ordered, k)
            reach = max(reach, ordered[k][2][hi])
    if best is None:
        return sorted(items, key=lambda item: (item[2][1], item[2][0]))
    _, ordered, k = best
    return _xy_cut(ordered[:k]) + _xy_cut(ordered[k:])


def _row_order(items) -> list:
    return sorted(items, key=lambda item: (round(item[2][1] / 10), item[2][0]))


def _clean_llm_text(text: str) -> str:
    text = re.sub(r'<[^>]+>', ' ', text)
    text = re.sub(r'^\s{0,3}#{1,6}\s+', '', text, flags=re.M)
    text = re.sub(r'^[\s|:\-]*-{3,}[\s|:\-]*$', '', text, flags=re.M)
    return text.replace('**', '').replace('|', ' ')


def _matcher_key(text: str) -> str:
    return ''.join(' ' if c.isspace() else (c.lower() if len(c.lower()) == 1 else c) for c in text)


def _target_order(items, key: str, gram: int = 4, max_hits: int = 12) -> list:
    from collections import Counter
    index: Dict[str, List[int]] = {}
    for j in range(len(key) - gram + 1):
        index.setdefault(key[j:j + gram], []).append(j)
    ordered = _xy_cut(items)
    estimates = []
    previous = 0.0
    for n, item in enumerate(ordered):
        line = _matcher_key(item[0])
        votes = Counter()
        grams = 0
        for i in range(len(line) - gram + 1):
            hits = index.get(line[i:i + gram], ())
            if 0 < len(hits) <= max_hits:
                grams += 1
                for j in hits:
                    votes[(j - i) // gram] += 1
        stripped = line.strip()
        if votes:
            offset, count = votes.most_common(1)[0]
            if count >= max(1 if grams <= 2 else 2, grams * 0.25):
                previous = offset * gram
                votes = None
        if votes is not None and stripped and key.count(stripped) == 1:
            previous = key.find(stripped)
        estimates.append((previous, n))
    return [ordered[n] for _, n in sorted(estimates)]


def _snap_cut(target: str, cut: int, window: int = 12) -> int:
    if cut <= 0 or cut >= len(target) or target[cut - 1].isspace() or target[cut].isspace():
        return cut
    for distance in range(1, window + 1):
        for candidate in (cut - distance, cut + distance):
            if 0 < candidate < len(target) and target[candidate].isspace():
                return candidate
    return cut


def _align_items(items, text: str):
    from difflib import SequenceMatcher
    target = _clean_llm_text(text)
    key = _matcher_key(target)
    best = None
    for ordered in (_target_order(items, key), _row_order(items), _xy_cut(items)):
        source = ' '.join(item[0] for item in ordered)
        matcher = SequenceMatcher(None, _matcher_key(source), key, autojunk=False)
        matched = sum(block.size for block in matcher.get_matching_blocks())
        ratio = 2 * matched / max(1, len(source) + len(key))
        if best is None or ratio > best[0]:
            best = (ratio, ordered, matcher.get_opcodes())
    ratio, ordered, opcodes = best
    if ratio < MIN_ALIGN_RATIO:
        return None

    def to_target(pos: int) -> int:
        for tag, i1, i2, j1, j2 in opcodes:
            if i1 <= pos < i2:
                if tag == 'equal':
                    return j1 + pos - i1
                return j1 + round((pos - i1) * (j2 - j1) / (i2 - i1))
        return len(target)

    cuts = [0]
    position = 0
    for item in ordered[:-1]:
        position += len(item[0])
        cuts.append(max(cuts[-1], _snap_cut(target, to_target(position))))
        position += 1
    cuts.append(len(target))

    aligned = []
    for n, item in enumerate(ordered):
        segment = ' '.join(target[cuts[n]:cuts[n + 1]].split())
        aligned.append((segment or item[0], item[1], item[2]))
    return aligned


def _paragraphs_from_boxes(items) -> str:
    items = sorted(items, key=lambda item: (round(item[2][1] / 10), item[2][0]))
    paragraphs: List[str] = []
    current: List[str] = []
    previous_bottom = previous_height = None
    for text, _, (x0, y0, x1, y1) in items:
        height = y1 - y0
        if previous_bottom is not None and y0 - previous_bottom > max(height, previous_height) * 0.9:
            paragraphs.append('\n'.join(current))
            current = []
        current.append(text)
        previous_bottom, previous_height = y1, height
    if current:
        paragraphs.append('\n'.join(current))
    return '\n\n'.join(paragraphs)


class OCRService(BaseConverter):
    ENGINES = ('rapid', 'qwen')
    DEFAULT_ENGINE = 'rapid'

    _rapid = None
    _rapid_key = None
    _rapid_lock = threading.Lock()
    _status_cache = {'at': 0.0, 'data': None}

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.llm = LLMService()
        self.info: Dict[str, Any] = {}
        self.rapid_quality = DEFAULT_RAPID_QUALITY
        self.rapid_script = DEFAULT_RAPID_SCRIPT

    @classmethod
    def engine_status(cls, max_age: float = 20.0) -> Dict[str, Dict[str, Any]]:
        cache = cls._status_cache
        if cache['data'] is not None and time.time() - cache['at'] < max_age:
            return cache['data']

        llm = LLMService()
        status = {}

        try:
            import rapidocr
            status['rapid'] = {'available': True, 'reason': None}
        except ImportError:
            status['rapid'] = {'available': False, 'reason': 'missing_rapidocr'}

        ocr_model = os.environ.get('LLM_OCR_MODEL', 'qwen/qwen3-vl-8b')
        qwen = llm.model_status(ocr_model)
        status['qwen'] = {'available': qwen == 'ok', 'reason': None if qwen == 'ok' else qwen, 'model': ocr_model}


        formatting_model = os.environ.get('LLM_FORMATTING_MODEL', 'ministral-3-8b-instruct-2512')
        correction = llm.model_status(formatting_model)
        status['correction'] = {'available': correction == 'ok', 'reason': None if correction == 'ok' else correction, 'model': formatting_model}


        cache['at'], cache['data'] = time.time(), status
        return status

    def _resolve_engine(self, options: Dict[str, Any]) -> Tuple[str, bool]:
        engine = options.get('ocr_engine') or self.DEFAULT_ENGINE
        if engine not in self.ENGINES:
            raise OCRError(f"Unknown OCR engine: {engine}", code='err_ocr_engine')
        status = self.engine_status()
        info = status[engine]
        if not info['available']:
            raise OCRError(f"OCR engine '{engine}' is not available ({info['reason']})", code=f"err_ocr_{info['reason']}")
        correct = engine == 'rapid' and bool(options.get('use_llm'))
        if correct and not status['correction']['available']:
            raise OCRError("AI text correction is not available", code=f"err_ocr_{status['correction']['reason']}")
        quality = options.get('ocr_quality')
        self.rapid_quality = quality if quality in RAPID_QUALITIES else DEFAULT_RAPID_QUALITY
        script = options.get('ocr_script')
        self.rapid_script = script if script in RAPID_SCRIPTS else DEFAULT_RAPID_SCRIPT
        return engine, correct

    @staticmethod
    def _rapid_params(quality: str, script: str) -> Dict[str, Any]:
        from rapidocr import ModelType, OCRVersion, LangDet, LangRec
        profile = RAPID_QUALITIES[quality]
        params = {'Global.log_level': 'error', 'Global.max_side_len': profile['max_side']}
        if script == 'latin':
            params.update({
                'Det.ocr_version': OCRVersion.PPOCRV6, 'Det.model_type': ModelType(profile['model']),
                'Rec.ocr_version': OCRVersion.PPOCRV6, 'Rec.model_type': ModelType(profile['model']),
            })
        else:
            params.update({
                'Det.ocr_version': OCRVersion.PPOCRV5, 'Det.lang_type': LangDet.MULTI,
                'Det.model_type': ModelType.SERVER if quality == 'accurate' else ModelType.MOBILE,
                'Rec.ocr_version': OCRVersion.PPOCRV5, 'Rec.model_type': ModelType.MOBILE,
                'Rec.lang_type': LangRec(script),
            })
        return params

    @classmethod
    def _get_rapid(cls, quality: str = DEFAULT_RAPID_QUALITY, script: str = DEFAULT_RAPID_SCRIPT):
        key = (quality, script)
        with cls._rapid_lock:
            if cls._rapid_key != key:
                cls._release_rapid_locked()
                from rapidocr import RapidOCR
                cls._rapid = RapidOCR(params=cls._rapid_params(quality, script))
                cls._rapid_key = key
            return cls._rapid

    @classmethod
    def _release_rapid_locked(cls):
        if cls._rapid is not None:
            cls._rapid = None
            cls._rapid_key = None
            gc.collect()

    @classmethod
    def release_rapid(cls):
        with cls._rapid_lock:
            cls._release_rapid_locked()

    @property
    def rapid_dpi(self) -> int:
        return RAPID_QUALITIES[self.rapid_quality]['dpi']

    @staticmethod
    def _padded_image(image_path: str):
        from PIL import Image
        with Image.open(image_path) as source:
            image = source.convert('RGB')
        width, height = image.size
        if min(width, height) >= DET_MIN_SIDE:
            return image
        canvas = Image.new('RGB', (max(width, DET_MIN_SIDE), max(height, DET_MIN_SIDE)), image.getpixel((0, 0)))
        canvas.paste(image, (0, 0))
        return canvas

    def _rapid_raw(self, image_path: str, quality: Optional[str] = None, script: Optional[str] = None):
        engine = self._get_rapid(quality or self.rapid_quality, script or self.rapid_script)
        result = engine(self._padded_image(image_path))
        boxes = result.boxes if result.boxes is not None else []
        items = []
        for box, text, score in zip(boxes, result.txts or [], result.scores or []):
            text = (text or '').strip()
            if not text:
                continue
            xs = [float(p[0]) for p in box]
            ys = [float(p[1]) for p in box]
            items.append((text, float(score), (min(xs), min(ys), max(xs), max(ys))))
        return items

    @staticmethod
    def _orientation_quality(items) -> Tuple[float, float]:
        if not items:
            return 0.0, 0.0
        vertical = sum(1 for _, _, (x0, y0, x1, y1) in items if (y1 - y0) > (x1 - x0) * 1.5)
        quality = sum(len(text) * score for text, score, (x0, y0, x1, y1) in items if (x1 - x0) >= (y1 - y0))
        return quality, vertical / len(items)

    def _recognize_rapid(self, image_path: str) -> OcrPage:
        items = self._rapid_raw(image_path)
        return OcrPage(_paragraphs_from_boxes(items), [])

    @staticmethod
    def _page_lines(items, to_page, offset, zoom) -> list:
        import fitz
        lines = []
        for text, _, (x0, y0, x1, y1) in items:
            height = y1 - y0
            baseline = y1 - height * 0.18
            start = (fitz.Point(x0, baseline) + offset) * to_page
            end = (fitz.Point(x1, baseline) + offset) * to_page
            lines.append((text, start, end, height / zoom))
        return lines

    def _layout_items(self, image_path: str, text: str) -> list:
        script = _detect_script(text)
        try:
            return self._rapid_raw(image_path, LAYOUT_QUALITY if script == 'latin' else 'accurate', script)
        except Exception as e:
            logger.warning("OCR layout detection failed: %s", e)
            return []

    def _ocr_pdf_page(self, page, engine: str, output_format: str, options: Dict[str, Any], temp_path: str, with_layout: bool = False) -> OcrPage:
        import fitz
        zoom = (self.rapid_dpi if engine == 'rapid' else OCR_DPI) / 72.0
        if engine != 'rapid':
            matrix = fitz.Matrix(zoom, zoom)
            pix = page.get_pixmap(matrix=matrix)
            pix.save(temp_path)
            result = self.recognize(temp_path, engine, output_format, options)
            if with_layout and result.text.strip():
                items = self._layout_items(temp_path, result.text)
                if items:
                    result.layout = (items, ~matrix * page.derotation_matrix, fitz.Point(pix.x, pix.y), zoom)
            return result

        best = None
        for angle in (0, 90, 270):
            matrix = fitz.Matrix(zoom, zoom).prerotate(angle)
            pix = page.get_pixmap(matrix=matrix)
            pix.save(temp_path)
            items = self._rapid_raw(temp_path)
            quality, vertical_ratio = self._orientation_quality(items)
            if best is None or quality > best[0]:
                best = (quality, items, matrix, fitz.Point(pix.x, pix.y))
            if angle == 0 and (not items or vertical_ratio < 0.3):
                break

        _, items, matrix, offset = best
        to_page = ~matrix * page.derotation_matrix
        lines = self._page_lines(items, to_page, offset, zoom)
        return OcrPage(_paragraphs_from_boxes(items), lines, (items, to_page, offset, zoom) if items else None)

    def _recognize_llm(self, image_path: str, output_format: str, options: Dict[str, Any]) -> OcrPage:
        prompt = (
            "Transcribe the text from this image exactly as it appears. "
            "Output ONLY plain text. Do NOT use Markdown formatting, do NOT use HTML tags. "
            "Keep paragraphs separated by blank lines."
        )
        if output_format == 'md':
            prompt = (
                "Transcribe the content of this image into well-formatted Markdown. "
                "Use # and ## for headings, - for bullet lists, 1. for numbered lists, **bold**, "
                "| tables | with | pipes | and > for quotes. Output ONLY the raw markdown, no code fences."
            )
        elif output_format == 'html':
            theme = 'a DARK theme (background #1a1a1a, text #e0e0e0)' if options.get('ocr_theme') == 'dark' else 'a LIGHT theme (white background, dark text)'
            css_limit = ''
            if options.get('css_limit_enabled'):
                css_limit = f" Keep the CSS under about {int(options.get('css_limit_value', 1000))} characters."
            prompt = (
                "Transcribe the content of this image into clean, semantic HTML5 (h1, h2, p, ul, table...). "
                f"Add one small <style> block for {theme}.{css_limit} "
                "Do not embed images. Output ONLY the raw HTML."
            )
        try:
            text = self.llm.generate_from_image(image_path, prompt)
        except Exception as e:
            raise OCRError(str(e), code='err_ocr_llm_offline')
        return OcrPage(self._strip_fences(text))


    def recognize(self, image_path: str, engine: str, output_format: str = 'txt', options: Optional[Dict[str, Any]] = None) -> OcrPage:
        options = options or {}
        if engine == 'rapid':
            return self._recognize_rapid(image_path)
        return self._recognize_llm(image_path, output_format, options)

    @staticmethod
    def _strip_fences(text: str) -> str:
        text = (text or '').strip()
        text = re.sub(r'^```[a-zA-Z]*\s*\n?', '', text)
        text = re.sub(r'\n?```\s*$', '', text)
        return text.strip()

    def _iter_ocr_pages(self, input_path: str, engine: str, output_format: str, options: Dict[str, Any], temp_dir: str):
        import fitz
        if Path(input_path).suffix.lower() in IMAGE_EXTENSIONS:
            yield 0, 1, self.recognize(input_path, engine, output_format, options)
            return
        with fitz.open(input_path) as doc:
            if doc.needs_pass:
                raise OCRError("PDF is password protected", code='err_pdf_password')
            total = doc.page_count
            for i, page in enumerate(doc):
                if self.is_cancelled:
                    return
                yield i, total, self._ocr_pdf_page(page, engine, output_format, options, os.path.join(temp_dir, f'page_{i}.png'))

    def _engine_model(self, engine: str):
        return LLMService.ocr_model() if engine == 'qwen' else None

    def _correct_pages(self, texts: List[str], output_format: str, start: int, end: int) -> List[str]:
        corrected = []
        with LLMService.using(LLMService.formatting_model()):
            for i, text in enumerate(texts):
                if self.is_cancelled:
                    return []
                corrected.append(self.llm.correct_text(text, output_format) if text.strip() else text)
                self.report_progress(start + int((i + 1) / len(texts) * (end - start)))
        return corrected

    def _recognize_document(self, input_path: str, output_format: str, options: Dict[str, Any]) -> List[str]:
        engine, correct = self._resolve_engine(options)
        correct_pages = correct and output_format != 'html'
        ocr_end = 60 if correct_pages else 90
        temp_dir = tempfile.mkdtemp(prefix='ocr_')
        texts = []
        try:
            with LLMService.using(self._engine_model(engine)):
                for i, total, page in self._iter_ocr_pages(input_path, engine, output_format, options, temp_dir):
                    if self.is_cancelled:
                        return []
                    texts.append(page.text)
                    self.report_progress(5 + int((i + 1) / total * (ocr_end - 5)))
        finally:
            shutil.rmtree(temp_dir, ignore_errors=True)
        if correct_pages and texts:
            texts = self._correct_pages(texts, output_format, ocr_end, 90)
            if self.is_cancelled:
                return []
        self.info = {'engine': engine, 'correction': correct, 'pages': len(texts)}
        if engine == 'rapid':
            self.info.update({'quality': self.rapid_quality, 'script': self.rapid_script})
        if not any(t.strip() for t in texts):
            raise OCRError("No text was recognized", code='err_ocr_no_text')
        return texts

    def get_text_from_image(self, input_path: str, lang: Optional[str] = None, options: Optional[Dict[str, Any]] = None) -> str:
        options = dict(options or {})
        if lang and 'ocr_script' not in options:
            options['ocr_script'] = LANG_TO_SCRIPT.get(lang.lower().split('-')[0], DEFAULT_RAPID_SCRIPT)
        return '\n\n'.join(self._recognize_document(input_path, 'txt', options))

    def _html_document(self, pages: List[str], options: Dict[str, Any], engine_html: bool) -> str:
        dark = options.get('ocr_theme') == 'dark'
        if engine_html:
            styles = []
            bodies = []
            for page in pages:
                styles.extend(re.findall(r'<style[^>]*>(.*?)</style>', page, flags=re.S | re.I))
                body = re.search(r'<body[^>]*>(.*?)</body>', page, flags=re.S | re.I)
                content = body.group(1) if body else re.sub(r'</?(html|head|body)[^>]*>|<!DOCTYPE[^>]*>|<style[^>]*>.*?</style>|<title>.*?</title>|<meta[^>]*>', '', page, flags=re.S | re.I)
                bodies.append(f'<section class="page">{content.strip()}</section>')
            style = styles[0] if styles else ''
        else:
            bodies = []
            for page in pages:
                paragraphs = ''.join(f'<p>{html.escape(p).replace(chr(10), "<br>")}</p>' for p in page.split('\n\n') if p.strip())
                bodies.append(f'<section class="page">{paragraphs}</section>')
            style = ''
        base_style = (
            f"body{{font-family:system-ui,sans-serif;max-width:820px;margin:40px auto;padding:0 20px;line-height:1.6;"
            f"background:{'#1a1a1a' if dark else '#fff'};color:{'#e0e0e0' if dark else '#1a1a1a'}}}"
            f".page+.page{{border-top:1px solid {'#333' if dark else '#ddd'};margin-top:32px;padding-top:32px}}"
        )
        return (
            '<!DOCTYPE html><html><head><meta charset="utf-8"><title>OCR</title>'
            f'<style>{base_style}{style}</style></head><body>{"".join(bodies)}</body></html>'
        )

    def _write_docx(self, pages: List[str], output_path: str):
        from docx import Document
        from docx.enum.text import WD_BREAK
        document = Document()
        for n, page in enumerate(pages):
            for paragraph in page.split('\n\n'):
                if paragraph.strip():
                    document.add_paragraph(paragraph.strip())
            if n < len(pages) - 1:
                document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        document.save(output_path)

    def _make_searchable_pdf(self, input_path: str, output_path: str, options: Dict[str, Any]) -> str:
        import fitz
        engine, correct = self._resolve_engine(options)
        force = bool(options.get('ocr_force'))
        font = fitz.Font('helv')
        ext = Path(input_path).suffix.lower()

        if ext in IMAGE_EXTENSIONS:
            with fitz.open(input_path) as img:
                doc = fitz.open('pdf', img.convert_to_pdf())
        else:
            doc = fitz.open(input_path)
            if doc.needs_pass:
                doc.close()
                raise OCRError("PDF is password protected", code='err_pdf_password')

        temp_dir = tempfile.mkdtemp(prefix='ocr_pdf_')
        recognized = skipped = 0
        blocks = []
        try:
            total = doc.page_count
            ocr_end = 60 if correct else 95
            with LLMService.using(self._engine_model(engine)):
                for i, page in enumerate(doc):
                    if self.is_cancelled:
                        return None
                    if not force and page.get_text().strip():
                        skipped += 1
                        self.report_progress(5 + int((i + 1) / total * (ocr_end - 5)))
                        continue
                    result = self._ocr_pdf_page(page, engine, 'txt', options, os.path.join(temp_dir, f'page_{i}.png'), with_layout=True)
                    if result.lines and not correct:
                        self._place_lines(page, result.lines, font)
                    elif result.text.strip():
                        blocks.append((i, result.text, result.layout))
                    recognized += 1
                    self.report_progress(5 + int((i + 1) / total * (ocr_end - 5)))
            texts = [text for _, text, _ in blocks]
            if correct and texts:
                texts = self._correct_pages(texts, 'txt', ocr_end, 95)
                if self.is_cancelled:
                    return None
            positioned = 0
            for (index, _, layout), text in zip(blocks, texts):
                aligned = _align_items(layout[0], text) if layout else None
                if aligned:
                    self._place_lines(doc[index], self._page_lines(aligned, *layout[1:]), font)
                    positioned += 1
                else:
                    self._place_block(doc[index], text, font)
            doc.save(output_path, garbage=3, deflate=True)
        finally:
            doc.close()
            shutil.rmtree(temp_dir, ignore_errors=True)
        self.info = {'engine': engine, 'pages_ocr': recognized, 'pages_skipped': skipped}
        if blocks:
            self.info['pages_positioned'] = positioned
        if engine == 'rapid':
            self.info.update({'quality': self.rapid_quality, 'script': self.rapid_script})
        return output_path

    @staticmethod
    def _place_lines(page, lines, font):
        import math
        import fitz
        for text, start, end, height in lines:
            dx, dy = end.x - start.x, end.y - start.y
            width = math.hypot(dx, dy)
            if width <= 0 or height <= 0:
                continue
            fontsize = height * 0.85
            length = font.text_length(text, fontsize=fontsize)
            if length > 0:
                fontsize = max(height * 0.4, min(height * 1.3, fontsize * width / length))
            angle = math.degrees(math.atan2(dy, dx))
            writer = fitz.TextWriter(page.rect)
            writer.append(start, text, font=font, fontsize=fontsize)
            morph = (start, fitz.Matrix(-angle)) if abs(angle) > 0.5 else None
            writer.write_text(page, render_mode=3, morph=morph)

    @staticmethod
    def _place_block(page, text, font):
        import fitz
        rect = fitz.Rect(page.rect * page.derotation_matrix).normalize() + (24, 24, -24, -24)
        for fontsize in (10, 8, 6, 4, 3):
            writer = fitz.TextWriter(page.rect)
            overflow = writer.fill_textbox(rect, text, font=font, fontsize=fontsize)
            if not overflow or fontsize == 3:
                writer.write_text(page, render_mode=3)
                return

    def convert(self, input_path: str, output_path: str, output_format: str, options: Optional[Dict[str, Any]] = None) -> str:
        options = options or {}
        output_format = output_format.lower()
        if output_format.startswith('ocr-'):
            output_format = output_format[4:]
        if output_format == 'pdf_ocr':
            output_format = 'pdf'

        self.report_progress(3)
        if output_format == 'pdf':
            result = self._make_searchable_pdf(input_path, output_path, options)
            self.report_progress(100)
            return result

        pages = self._recognize_document(input_path, output_format, options)
        if self.is_cancelled:
            return None
        engine, correct = self.info.get('engine'), self.info.get('correction')

        if output_format == 'docx':
            self._write_docx(pages, output_path)
        elif output_format == 'html':
            if engine in ('qwen',):
                content = self._html_document(pages, options, engine_html=True)
            elif correct:
                css_limit = int(options.get('css_limit_value', 1000)) if options.get('css_limit_enabled') else None
                content = self._strip_fences(self.llm.generate_html('\n\n'.join(pages), options.get('ocr_theme', 'light'), css_limit))
            else:
                content = self._html_document(pages, options, engine_html=False)
            with open(output_path, 'w', encoding='utf-8') as f:
                f.write(content)
        else:
            separator = '\n\n---\n\n' if output_format == 'md' else '\n\n'
            with open(output_path, 'w', encoding='utf-8') as f:
                f.write(separator.join(pages))

        self.report_progress(100)
        return output_path

    @staticmethod
    def get_supported_input_formats() -> set:
        return {'png', 'jpg', 'jpeg', 'webp', 'bmp', 'tiff', 'tif', 'pdf'}

    @staticmethod
    def get_supported_output_formats() -> set:
        return {'txt', 'md', 'html', 'ocr-txt', 'ocr-md', 'ocr-html', 'docx', 'ocr-docx', 'pdf', 'ocr-pdf'}
